import json, pickle, os, numpy as np
EV = json.load(open('evidence/onchain_evidence.json')); TXS = '; '.join(e['tx_hash'][:12] + '...' for e in EV)
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
exec(open("build_notebook.py").read().split("cells = []")[0])        # reuse the same computed numbers N
NAVY, TEAL, GREY, LG, RED, GREEN, AMBER, INK = [RGBColor.from_string(x) for x in ("0B1F3A", "0A9396", "6B7785", "F1F4F8", "C0392B", "2E8B57", "E09F3E", "1B2430")]
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5); BL = prs.slide_layouts[6]; page = [0]
def tb(s, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Arial"):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(6)
        r = p.add_run(); r.text = ln; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return b
def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); r.fill.solid(); r.fill.fore_color.rgb = fill
    if line: r.line.color.rgb = line
    else: r.line.fill.background()
    r.shadow.inherit = False; return r
def slide(title, tracker, source=None, notes=""):
    s = prs.slides.add_slide(BL); page[0] += 1
    tb(s, .55, .22, 8, .3, tracker.upper(), 10, True, TEAL)
    tb(s, .55, .5, 12.2, 1.0, title, 24, True, NAVY, anchor=MSO_ANCHOR.TOP); rect(s, .6, 1.5, 12.1, .03, TEAL)
    if source: tb(s, .55, 7.0, 11.3, .3, "Source: " + source, 9, False, GREY)
    tb(s, 12.2, 7.0, .6, .3, str(page[0]), 10, False, GREY, PP_ALIGN.RIGHT); tb(s, 9.3, .22, 3.5, .3, "NULL Wallet", 10, True, GREY, PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes; return s
def img(s, name, x, y, w=None, h=None):
    kw = {}; 
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return s.shapes.add_picture(f"figs/{name}.png", Inches(x), Inches(y), **kw)
def callout(s, text, x=.6, y=6.25, w=12.1, h=.62, fill=NAVY):
    rect(s, x, y, w, h, fill); tb(s, x + .15, y, w - .3, h, text, 14, True, RGBColor(255, 255, 255), anchor=MSO_ANCHOR.MIDDLE)
def bullets(s, items, x, y, w, size=15, gap=.1):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(4.5)); tf = b.text_frame; tf.word_wrap = True
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(10)
        r = p.add_run(); r.text = head + " "; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY; r.font.name = "Arial"
        r2 = p.add_run(); r2.text = body; r2.font.size = Pt(size); r2.font.color.rgb = INK; r2.font.name = "Arial"
def stat(s, x, y, big, small, color=TEAL, w=2.9):
    rect(s, x, y, w, 1.55, LG); tb(s, x, y + .1, w, .8, big, 34, True, color, PP_ALIGN.CENTER); tb(s, x + .1, y + .9, w - .2, .65, small, 11.5, False, INK, PP_ALIGN.CENTER)

# 1 title
s = prs.slides.add_slide(BL); page[0] += 1; rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, .8, 3.55, 1.6, .06, TEAL)
tb(s, .8, 1.6, 11.5, 1.0, "NULL Wallet", 54, True, RGBColor(255, 255, 255)); tb(s, .8, 2.65, 11.5, 0.9, "The AI wallet that must reject \"no edge\" before it moves money", 26, False, RGBColor(220, 232, 240))
tb(s, .8, 3.8, 11.5, 1.2, ["2026 Shenzhen International FinTech Competition  |  International Track  |  Topic F: Agent Wallet", "Theme: AI as a Financial Participant"], 15, False, RGBColor(170, 190, 205))
tb(s, .8, 5.9, 11.5, 0.8, "Team: AskWallet   |   Brandon Haris Effendi (team leader, MSc Finance), Maulik Bansal (MSc Finance), Desylionie Onggani Winata (MSc Logistics Management), Lund University, Sweden   |   br5880ha-s@student.lu.se", 14, True, RGBColor(255, 255, 255))
s.notes_slide.notes_text_frame.text = "(20 s) We built a wallet for AI agents that has to earn the right to move money. Before it trades, it must statistically reject the hypothesis that it has no edge. Everything it decides, including every refusal, is recorded and can be verified."

# 2 exec summary
s = slide("AI agents with wallets need proof, not just permission: NULL Wallet moves money only after rejecting \"no edge\"", "Executive summary", "Team analysis; Monte Carlo of %s simulated markets; replay of real BTC hourly data %s to %s" % (N["h0m"], N["periodstart"], N["periodend"]),
          "(45 s) Answer first. Today's agent wallets limit how much an agent can spend. They do not check whether the reason is real. We add an Evidence Court built from our three assignments. In 1,000 no-edge markets it never approved a trade. On real BTC it refused all 36 monthly decisions and so avoided the 20 percent loss the same models suffered without it.")
bullets(s, [("The gap.", "Spend limits and allowlists ask \"is this transfer permitted?\" Nobody asks \"is the reason real?\" An agent can obey every limit and still trade on noise."),
            ("Our answer.", "A wallet that must reject H0 (\"no net edge after costs\") with pre-registered, multiplicity-corrected tests before value moves; every refusal is recorded and anchored on-chain."),
            ("The proof.", "No-edge markets: %s%% false GO in %s runs. Real edge: GO rises to %s%% at moderate strength. Real BTC: 36 of 36 monthly verdicts say do not trade." % (N["h0go"], N["h0m"], N["p20"])),
            ("The safety.", "No language model sits between a request and the key: strict intents, tiers, caps, allowlist, kill-switch. %s automated tests incl. prompt-injection attacks." % N["ntests"])], .6, 1.75, 7.4, 14.5)
stat(s, 8.4, 1.8, N["h0go"] + "%", "false GO in %s no-edge simulated markets" % N["h0m"], GREEN, 2.1); stat(s, 10.7, 1.8, N["p20"] + "%", "GO when a moderate edge truly exists", TEAL, 2.1)
stat(s, 8.4, 3.55, "0 of 36", "monthly verdicts on real BTC allowed trading", NAVY, 2.1); stat(s, 10.7, 3.55, "-20%", "wallet loss of the same models run without the court", RED, 2.1)
callout(s, "So what: evidence before permission makes an AI agent's wallet auditable, and refusing to trade becomes a result you can prove.")

# 3 fishbone
s = slide("Six families of causes make agents trade on false edges, and each one has a control in NULL Wallet", "Problem", "Team root-cause analysis (fishbone); controls map to code modules in the repository",
          "(50 s) Why do agents lose money on patterns that are not real? Data leakage, overfitting and trying many strategies, ignored costs, regime shocks, execution attacks and weak governance. For every cause the teal text names the control we built.")
img(s, "05_fishbone", .6, 1.65, w=12.1)
# 4 architecture
s = slide("Two separated pipes decide whether to act and how to act, and meet only at the latest GO receipt", "Solution", "NULL Wallet architecture (nullwallet/ package)",
          "(50 s) Top pipe: data to causal features to walk-forward models to the Evidence Court, which writes a receipt into a hash-chained ledger. Bottom pipe: a request becomes a strict typed intent, then a deterministic policy engine, then the isolated signer. The bridge is the latest GO receipt.")
img(s, "01_architecture", .6, 1.65, w=8.3)
bullets(s, [("Evidence pipe.", "Decides whether the strategy deserves capital."), ("Action pipe.", "Decides whether this specific move is allowed. Plain code, no model."), ("Bridge.", "A tier-2 move must cite the latest, fresh GO receipt."), ("Ledger.", "Hash-chained; head anchored on-chain.")], 9.1, 1.8, 3.7, 13)
callout(s, "So what: a compromised or confused language layer still cannot sign anything the policy engine does not approve.")
# 5 court
s = slide("The Evidence Court turns six standard econometric tests into one verdict: GO, WATCH or HOLD", "Solution", "Diebold-Mariano (1995), Newey-West (1987), Holm (1979), White (2000)",
          "(60 s) Each gate is something we already did in coursework. G0 no look-ahead. G1 assumption tests, which fail on real data so we always use robust errors. G2 skill against the prevailing mean with Diebold-Mariano and Holm over five candidates. G3 survive without shock bars. G4 net edge after costs with a bootstrap bound. G5 break-even cost at least twice the cost. All pass is GO; skill without net edge is WATCH; else HOLD, the default.")
img(s, "02_gates", .6, 1.65, w=12.1)
callout(s, "So what: \"no trade\" is the default and every GO has to survive data-snooping control and realistic costs.", y=6.3)
# 6 security tiers
s = slide("Authority grows only as evidence grows: four tiers, and anything unknown is refused", "Security", "NULL Wallet policy engine (nullwallet/policy.py); tests in tests/test_redteam.py",
          "(45 s) T0 observe. T1 attest on-chain, which can cost only gas. T2 move value, which needs a fresh GO receipt, a whitelisted vault, per-transaction and daily caps and a cooldown. T3 is everything else and is always refused. A drawdown kill-switch needs a human to reset.")
img(s, "03_tiers", .6, 1.65, w=7.2)
bullets(s, [("Strict intent schema.", "Unknown fields are rejected."), ("Exact allowlist.", "Look-alike addresses fail."), ("Freshness and cooldown.", "Old approvals cannot be replayed."), ("Isolated signer.", "Key only in an environment variable, never in logs or receipts."), ("Kill-switch.", "10% drawdown halts value moves.")], 8.1, 1.8, 4.7, 13)
# 7 bowtie + tests
s = slide("Five threats are blocked before the event and four consequences are limited after it, backed by %s tests" % N["ntests"], "Security", "Team bow-tie risk analysis; unittest results (all passing)",
          "(45 s) The bow-tie: on the left the threats and the barrier for each; on the right the consequences and how we limit them. We tested attacks, not just features: prompt injection, address poisoning, replay, a thousand requests in one bar, a tampered ledger and key leakage.")
img(s, "06_bowtie", .6, 1.6, w=8.2); img(s, "18_tests", 8.9, 1.75, w=3.9)
callout(s, "So what: even when the text parser is fooled, the policy engine still refuses (tested).", y=6.3)
# 8 monte carlo
s = slide("In %s no-edge markets the court never approved a trade, yet it approves %s%% when a moderate edge truly exists" % (N["h0m"], N["p20"]), "Evidence 1: stress test", "Monte Carlo (%s); GARCH volatility, Student-t shocks, jumps; planted mean-reversion edge b" % N["mcfile"],
          "(60 s) We tested the court the way one tests any estimator. One thousand simulated markets with no edge at all: zero false GOs, even when trading costs are set to zero. Then we plant an edge of increasing strength. GO rises from under one percent to about 92 percent at b equals 0.2 and 100 percent at b equals 0.4. It is strict, not blind.")
img(s, "17_power", .6, 1.65, w=7.0)
bullets(s, [("Size.", "%s%% GO under H0 (%s markets); %s%% with zero costs." % (N["h0go"], N["h0m"], N["h0nc"])), ("Power.", "%s%% (weak), %s%%, %s%% and %s%% as the edge strengthens." % (N["p05"], N["p10"], N["p20"], N["p40"])), ("Why trust it.", "Four gates must all pass; Holm correction prices in five tries."), ("Versus a naive picker.", "Any-of-5 with raw p<0.05 and no cost margin trades in %s%% of no-edge markets and %s%% at weak edge: more sensitive, but uncontrolled." % (N["naive"], N["naive05"])), ("Caveat.", "Per-month tests repeat; sequential control is on the roadmap.")], 7.9, 1.8, 4.9, 13)
callout(s, "So what: strict where there is nothing to find, sensitive where there is something, with the trade-off shown openly.", y=6.3)
# 9 real data
s = slide("On real BTC the court refused all 36 monthly decisions and avoided the 20% loss the same models suffered", "Evidence 2: real data", "BTC 5-min data aggregated to hourly, replay %s to %s; sandbox with assumed costs, not live trading" % (N["periodstart"], N["periodend"]),
          "(60 s) Same models, same data, same costs. Without the court the two agents traded constantly, paid about 2,400 to 2,800 dollars in costs on a 2,000 dollar active bucket and lost 20 percent of the wallet. The court said WATCH or HOLD every month, so it paid nothing. Buy and hold ended higher but with a 77 percent drawdown, and it is fully invested, so it is not like for like. This is a sandbox with assumed costs.")
img(s, "10_equity", .6, 1.65, w=7.6)
rect(s, 8.5, 1.8, 4.3, 3.9, LG); tb(s, 8.65, 1.85, 4.0, .4, "Final value of a $10,000 wallet", 13, True, NAVY)
rows = [("NULL Wallet", "$" + N["w_final"], TEAL), ("No court, GBM", "$" + N["g_final"], RED), ("No court, logit", "$" + N["l_final"], RED), ("Buy & hold*", "$" + N["bh_final"], GREY)]
for i, (a, b_, c) in enumerate(rows): tb(s, 8.65, 2.4 + i * .55, 2.4, .4, a, 14, False, INK); tb(s, 10.7, 2.4 + i * .55, 1.9, .4, b_, 16, True, c, PP_ALIGN.RIGHT)
tb(s, 8.65, 4.7, 4.0, 1.0, "*Fully invested, max drawdown %s%%. Not like-for-like: the court never claims to beat a bull market." % N["bh_dd"], 10.5, False, GREY)
callout(s, "So what: the court's value is not paying costs for an edge that does not survive them.", y=6.2)
# 10 why
s = slide("The models do beat a naive forecast, but their edge never reaches twice the cost: a result we reproduced", "Evidence 2: real data", "Court receipts on real BTC; modelled cost %s bp per switch (fee, slippage, square-root impact, gas), 2x margin required" % N["unit"],
          "(50 s) Of 180 candidate-window tests, 93 show statistical skill. But after costs none passes the net-edge gate. The median break-even cost of the best model is about 4.6 basis points against a modelled 7.5. Statistical edge, no economic edge, exactly what we found on real BTC.")
img(s, "13_funnel", .6, 1.65, w=6.1); img(s, "14_breakeven", 6.8, 1.65, w=6.0)
callout(s, "So what: WATCH is a real result: the pattern is statistically there but not tradable, so the wallet stays in reserve.", y=6.2)
# 11 failure found
s = slide("Stress-testing exposed a weakness that we fixed and disclose: a GO is evidence about the past, not a guarantee", "Evidence 3: honest limits", "Cost sweep on real BTC; regime guard added after the first replay (post hoc, therefore not an out-of-sample claim)",
          "(60 s) We kept attacking our own system. If costs are set to zero the court says GO in 10 windows and our first replay lost money in the March 2020 crash. We added a regime guard that flattens when recent volatility exceeds three times the evidence window. We label it post hoc. At a quarter of the base cost the court says GO once and that trade lost about five percent. So a GO is evidence about the past, and caps and the kill-switch matter.")
img(s, "15_cost_sweep", .6, 1.65, w=7.6)
bullets(s, [("Cost sweep.", "GO in %s of 36 windows at zero cost (8 on an independent rerun), %s at 0.25x, none at 0.5x or above." % (N["swgo0"], N["swgo25"])), ("A failure we found.", "Zero-cost run lost money in the March 2020 crash."), ("The fix.", "Regime guard: flat if recent volatility exceeds 3x the evidence window. Added after the fact, so we claim no out-of-sample benefit."), ("Still true.", "At 0.25x cost one GO lost about 5%: GO is not a guarantee.")], 8.4, 1.8, 4.4, 12.5)
callout(s, "So what: caps, kill-switch and guard are essential; the honesty of the receipt is the product.", y=6.3)
# 12 innovation
s = slide("What is new: evidence before permission, pre-registered and attested, with refusals as first-class records", "Innovation", "Public descriptions of agent-wallet products and arXiv 2601.04583 (survey); positioning is our qualitative judgement; we did not find the same combination",
          "(55 s) Existing wallets gate risk; we gate evidence. Five ideas: court gates from econometrics, data-snooping control, pre-registration by commit before reveal, attested refusals anchored on-chain, and a deterministic gate between language and signer. We did not find a system combining these; we cannot prove none exists.")
img(s, "08_positioning", .6, 1.65, w=5.6)
bullets(s, [("1. Court gates.", "Classical econometrics as a permission layer."), ("2. Data-snooping control.", "Holm over every candidate tried."), ("3. Pre-registration.", "Protocol hashed before the data is examined."), ("4. Attested refusals.", "Every HOLD and WATCH is a receipt, anchored on-chain."), ("5. Language never reaches the key.", "Typed intents and a deterministic policy gate.")], 6.5, 1.8, 6.3, 14)
callout(s, "So what: the standard tests are old; putting them in front of an agent's key, with proof of refusal, is the new part.", y=6.3)
# 13 coursework mapping
s = slide("Every gate is a standard, published method, so we can defend each line", "Innovation", "Standard methods", "(35 s) A quick map from methods to the six gates. This is why we can explain and defend every number.")
tbl = s.shapes.add_table(5, 3, Inches(.6), Inches(1.8), Inches(12.1), Inches(3.8)).table
data = [("Topic", "Method", "Where it lives in NULL Wallet"), ("Assumptions", "Assumption tests (Breusch-Pagan, Jarque-Bera, Ljung-Box), endogeneity, Monte Carlo of size and power", "G1 assumption audit; G0 look-ahead audit; the 1,000-market stress test"),
        ("Forecasting", "Pseudo-out-of-sample forecasts, Diebold-Mariano, ADF/KPSS, COVID dummies", "G2 skill duel; G3 shock robustness; walk-forward models"), ("Costs and multiple testing", "Net-of-cost edge, break-even cost, data snooping, neutral zone, HOLD/WATCH", "G4 net edge; G5 cost margin; Holm over candidates; dead-band; verdicts"), ("New", "Policy tiers, pre-registration, ledger, on-chain attestation", "Security and execution layers")]
for i, row in enumerate(data):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.text = v; p = c.text_frame.paragraphs[0]; p.runs[0].font.size = Pt(13); p.runs[0].font.name = "Arial"; p.runs[0].font.bold = (i == 0 or j == 0)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255) if i == 0 else INK; c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (LG if i % 2 == 0 else RGBColor(255, 255, 255))
tbl.columns[0].width = Inches(1.5); tbl.columns[1].width = Inches(5.6); tbl.columns[2].width = Inches(5.0)
# 14 on-chain + next
s = slide("Execution evidence is reproducible in three commands, and the roadmap moves from sandbox to live paper trading", "Execution and next steps", "Repository README and docs; on-chain transaction hashes are produced by the team on Base mainnet and listed in evidence/onchain_evidence.json",
          "(45 s) Judges can reproduce everything: install, run the tests, run the replay, open the web demo. The on-chain step runs on Base mainnet, a real blockchain, using a few USD of ETH from a throwaway wallet. It runs safety checks first, then writes zero-value attestations whose calldata is NULL1 plus the receipt hash. A capped vault transfer happens only if a GO exists and is explicitly allowed. Then the roadmap.")
rect(s, .6, 1.8, 6.1, 4.2, LG); tb(s, .8, 1.9, 5.8, .4, "Reproduce", 15, True, NAVY)
tb(s, .8, 2.4, 5.8, 3.4, ["pip install -r requirements.txt", "python -m unittest discover -s tests", "python run_all.py <ohlcv.csv>", "python -m nullwallet.webapp ledger.jsonl", "python onchain_demo.py ledger.jsonl --send"], 13, False, INK, font="Courier New")
tb(s, .8, 5.0, 5.8, .9, "On-chain on Base mainnet (real network, a few USD of ETH, throwaway wallet): attestations of the ledger head and refusals (T1); a capped vault transfer only if a GO exists (T2). Six mined transactions (status 1): " + TXS + "", 11, False, GREY)
img(s, "09_roadmap", 6.9, 1.9, w=5.9)
callout(s, "Ask: feedback on cost assumptions and a path to live paper trading with licensed data.", y=6.3, fill=TEAL)
# appendix
s = slide("Appendix A: the hash-chained ledger makes edits, deletions and reordering detectable", "Appendix", "nullwallet/ledger.py; tests/test_redteam.py (LedgerTamper)", "Backup slide.")
img(s, "04_ledger_chain", .6, 1.8, w=12.1)
s = slide("Appendix B: a language request can never reach the signer without passing the deterministic policy", "Appendix", "nullwallet/chat.py, policy.py, onchain.py", "Backup slide.")
img(s, "07_sequence", 1.4, 1.65, w=10.4)
s = slide("Appendix C: classical regression assumptions fail on real BTC returns, so the court always uses robust inference", "Appendix", "Gate G1 on the real BTC sample", "Backup slide.")
img(s, "19_assumptions", 1.2, 1.8, w=10.8)
s = slide("Appendix D: after costs, the mean excess return of the GBM and reversal models is below zero in 36 and 35 of 36 windows", "Appendix", "95% block-bootstrap intervals per 30-day window", "Backup slide.")
img(s, "16_excess_ci", .6, 1.8, w=12.1)
s = slide("Appendix E: known risks and limits we state openly", "Appendix", "docs/Security_Self_Assessment and technical documentation, Section 10", "Backup slide.")
lim = [("Sandbox, not live trading; costs are assumptions", "Cost sweep; live paper trading next"), ("A GO is evidence about the past", "Regime guard (post hoc), caps, kill-switch"), ("Monthly tests repeat, inflating false positives over time", "Four-gate rule; alpha-spending planned"), ("One main real dataset", "Synthetic markets; downloader for public data"),
       ("Ledger tail can be cut unless anchored", "On-chain head attestation"), ("Pure-Python signer is not constant-time", "Throwaway low-balance keys only; hardware/MPC signer for production"), ("Parser patterns can be evaded", "Policy engine is the real barrier")]
tbl = s.shapes.add_table(len(lim) + 1, 2, Inches(.6), Inches(1.8), Inches(12.1), Inches(4.5)).table
for i, row in enumerate([("Limitation", "Mitigation or plan")] + lim):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.text = v; r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(13); r.font.name = "Arial"; r.font.bold = i == 0; r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else INK
        c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (LG if i % 2 == 0 else RGBColor(255, 255, 255))
s = slide("Appendix F: six real Base mainnet transactions are mined, and each anchors a hash that exists in the ledger", "Appendix", "evidence/onchain_evidence.json; verify_onchain.py; chain id 8453", "Backup slide.")
bullets(s, [(e["label"] + ".", "block %s, status %s, %s" % (e["receipt"]["block"], e["receipt"]["status"], e["tx_hash"])) for e in EV], .8, 1.9, 11.8, size=13)
s = slide("Appendix G: our tests equal reference code to rounding error; the HAC test is slightly oversized, so G4 adds further gates", "Appendix", "audit_stats.py; results/stats_audit.json", "Backup slide.")
_SA = json.load(open("results/stats_audit.json"))["size_power"]
bullets(s, [(k.split(":")[0] + ".", "%.3f" % v) for k, v in list(_SA.items())[:8]], .8, 1.9, 11.8, size=13)
os.makedirs("deliverables", exist_ok=True); prs.save("deliverables/NullWallet_Deck.pptx"); print("slides", len(prs.slides))
